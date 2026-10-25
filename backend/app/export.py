"""DOCX export of a tailored resume (ATS-friendly: single column, real text, no tables or text boxes).

Mirrors the PDF layout: centred header, accent section headings with a hairline rule, and dates
right-aligned on the same line as the role via a right tab stop.
"""

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .schemas import TailoredResume

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x64, 0x64, 0x64)
ACCENT = RGBColor(0xBE, 0x3A, 0x26)
BODY_FONT = "Calibri"
NAME_FONT = "Georgia"
MARGIN = Mm(17)
CONTENT_W = Mm(210) - 2 * MARGIN


def _run(p, text: str, *, size: float = 10.5, bold: bool = False, italic: bool = False, color=INK, font: str | None = None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    if font:
        r.font.name = font
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    return r


def _para(doc: Document, *, before: float = 0, after: float = 0, align=None):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    if align is not None:
        p.alignment = align
    return p


def _bottom_border(p, color: str = "CDCDCD") -> None:
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "4", "w:space": "2", "w:color": color}.items():
        bottom.set(qn(k), v)
    borders.append(bottom)
    pPr.append(borders)


def _heading(doc: Document, text: str) -> None:
    p = _para(doc, before=11, after=4)
    p.paragraph_format.keep_with_next = True
    r = _run(p, text.upper(), size=9, bold=True, color=ACCENT)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "28")  # letter-spacing in twentieths of a point
    r._element.get_or_add_rPr().append(spacing)
    _bottom_border(p)


def _title_row(doc: Document, bold: str, rest: str = "", right: str = "", before: float = 5) -> None:
    p = _para(doc, before=before, after=1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(CONTENT_W, WD_TAB_ALIGNMENT.RIGHT)
    _run(p, bold, size=11, bold=True)
    if rest:
        _run(p, f"  ·  {rest}", size=10.5, color=MUTED)
    if right:
        _run(p, f"\t{right}", size=9.5, color=MUTED)


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(0), Pt(2)
    fmt.left_indent, fmt.first_line_indent = Mm(5), Mm(-4)
    _run(p, text)


def build_docx(r: TailoredResume) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.1
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)

    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = MARGIN
    section.top_margin = section.bottom_margin = Mm(15)

    _run(_para(doc, align=WD_ALIGN_PARAGRAPH.CENTER), r.name or "Resume", size=24, font=NAME_FONT)
    if r.headline:
        _run(_para(doc, before=2, align=WD_ALIGN_PARAGRAPH.CENTER), r.headline, size=11, color=MUTED)
    if r.contact:
        _run(_para(doc, before=3, align=WD_ALIGN_PARAGRAPH.CENTER), "   |   ".join(r.contact), size=9, color=MUTED)

    if r.summary:
        _heading(doc, "Summary")
        _run(_para(doc), r.summary)

    if r.skills:
        _heading(doc, "Skills")
        _run(_para(doc), "  ·  ".join(s.name for s in r.skills))

    if r.experience:
        _heading(doc, "Experience")
        for i, e in enumerate(r.experience):
            _title_row(doc, e.title, e.company, " · ".join(x for x in [e.location, e.dates] if x), before=0 if i == 0 else 7)
            for b in e.bullets:
                _bullet(doc, b)

    if r.projects:
        _heading(doc, "Projects")
        for i, pr in enumerate(r.projects):
            _title_row(doc, pr.name, before=0 if i == 0 else 6)
            if pr.details:
                p = _para(doc, after=1)
                p.paragraph_format.keep_with_next = True
                _run(p, pr.details, italic=True, color=MUTED)
            for b in pr.bullets:
                _bullet(doc, b)

    if r.education:
        _heading(doc, "Education")
        for i, ed in enumerate(r.education):
            _title_row(doc, ed.degree, ed.institution, ed.dates, before=0 if i == 0 else 4)
            if ed.details:
                _run(_para(doc), ed.details, size=10, color=MUTED)

    if r.certifications:
        _heading(doc, "Certifications")
        for c in r.certifications:
            _bullet(doc, c)

    for s in r.other_sections:
        if s.items:
            _heading(doc, s.heading)
            for item in s.items:
                _bullet(doc, item)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
